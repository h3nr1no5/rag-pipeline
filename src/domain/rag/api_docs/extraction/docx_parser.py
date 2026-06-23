"""DOCX parser for extracting raw structures from API documentation DOCX files."""

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


@dataclass
class RawTable:
    """Raw table extracted from a DOCX file."""

    headers: list[str]
    rows: list[list[str]]
    caption: str = ""
    position: int = -1


@dataclass
class RawParagraph:
    """Raw paragraph extracted from a DOCX file."""

    text: str
    style_name: str = ""
    heading_level: int = 0  # 0 for non-heading, 1-6 for heading levels
    position: int = -1


@dataclass
class RawDocument:
    """Raw document containing tables and paragraphs extracted from a DOCX."""

    tables: list[RawTable] = field(default_factory=list)
    paragraphs: list[RawParagraph] = field(default_factory=list)
    filename: str = ""


# ---------------------------------------------------------------------------
# Heading level detection
# ---------------------------------------------------------------------------

_HEADING_PREFIXES = [
    "heading",
    "head",
    "title",
    "caption",
]


def get_heading_level(style_name: str) -> int:
    """Extract heading level from a paragraph style name.

    Handles styles like "Heading 1", "Heading 2", "heading 1", "heading1",
    "Title", "Subtitle", etc.  Returns 0 if the style is not a heading.
    """
    if not style_name:
        return 0
    normalized = style_name.lower().replace(" ", "").replace("-", "")

    for prefix in _HEADING_PREFIXES:
        if normalized.startswith(prefix):
            rest = normalized[len(prefix) :]
            if rest.isdigit():
                level = int(rest)
                if 1 <= level <= 6:
                    return level

    # "Title" style → level 1
    if normalized == "title":
        return 1
    # "Subtitle" style → level 2
    if normalized == "subtitle":
        return 2

    return 0


# ---------------------------------------------------------------------------
# DocxParser
# ---------------------------------------------------------------------------


class DocxParser:
    """Parses DOCX files into raw document structures.

    Uses ``python-docx`` to open a ``.docx`` file and extract all tables
    (headers + rows) and paragraphs (text + heading level).  Tables and
    paragraphs are extracted in body order and each carries a ``position``
    field that reflects its index in the document body element stream.
    """

    def __init__(self, file_path: str | Path) -> None:
        self.file_path = Path(file_path)

    def parse(self) -> RawDocument:
        """Parse the DOCX file and return a :class:`RawDocument`.

        Raises:
            ValueError: If the file is corrupted or cannot be parsed.
        """
        try:
            doc = Document(str(self.file_path))
        except Exception as exc:
            raise ValueError(
                f"Failed to open or parse DOCX file '{self.file_path}': {exc}"
            ) from exc

        tables: list[RawTable] = []
        paragraphs: list[RawParagraph] = []

        table_idx = 0
        para_idx = 0
        position = 0

        body = doc.element.body
        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "tbl" and table_idx < len(doc.tables):
                raw_table = _extract_table(doc.tables[table_idx])
                raw_table.position = position
                tables.append(raw_table)
                table_idx += 1
                position += 1
            elif tag == "p" and para_idx < len(doc.paragraphs):
                raw_para = _extract_paragraph(doc.paragraphs[para_idx])
                raw_para.position = position
                paragraphs.append(raw_para)
                para_idx += 1
                position += 1
            else:
                # Other elements (sectPr, drawings, etc.) still occupy a position
                position += 1

        return RawDocument(
            tables=tables,
            paragraphs=paragraphs,
            filename=self.file_path.name,
        )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _extract_table(table) -> RawTable:
    """Extract a :class:`RawTable` from a ``python-docx`` Table object."""
    rows_data: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows_data.append(cells)

    if not rows_data:
        return RawTable(headers=[], rows=[], caption="")

    headers = rows_data[0]
    rows = rows_data[1:]
    return RawTable(headers=headers, rows=rows, caption="")


def _extract_paragraph(para: Paragraph) -> RawParagraph:
    """Extract a :class:`RawParagraph` from a ``python-docx`` Paragraph."""
    text = para.text.strip()
    style_name = para.style.name if para.style else ""
    heading_level = get_heading_level(style_name)
    return RawParagraph(
        text=text,
        style_name=style_name,
        heading_level=heading_level,
    )
