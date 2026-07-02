import asyncio
import logging
from abc import ABC, abstractmethod
from dataclasses import dataclass
from typing import Literal

logger = logging.getLogger(__name__)


@dataclass
class LinkInfo:
    """Represents a hyperlink extracted from a document."""
    type: Literal["internal", "external"]
    source_page: int | None = None
    target_page: int | None = None
    uri: str | None = None
    named_dest: str | None = None
    bbox: tuple[float, float, float, float] | None = None
    anchor_text: str | None = None


class DocumentParser(ABC):
    @abstractmethod
    async def parse(self, file_path: str) -> str:
        pass

    @abstractmethod
    def get_supported_extensions(self) -> list[str]:
        pass

    async def extract_links(self, file_path: str) -> list[LinkInfo]:
        """Extract hyperlinks from a document. Default returns empty list."""
        return []


class PDFParser(DocumentParser):
    def get_supported_extensions(self) -> list[str]:
        return [".pdf"]

    async def parse(self, file_path: str) -> str:
        def _parse():
            import fitz
            doc = fitz.open(file_path)
            text_parts = []

            for page_num, page in enumerate(doc):
                text = page.get_text()
                if text.strip():
                    text_parts.append(f"[Page {page_num + 1}]\n{text}")

            doc.close()
            return "\n\n".join(text_parts)

        return await asyncio.to_thread(_parse)

    async def extract_links(self, file_path: str) -> list[LinkInfo]:
        def _extract():
            import fitz

            doc = fitz.open(file_path)
            links: list[LinkInfo] = []

            for page_num, page in enumerate(doc):
                page_links = page.get_links()
                for link in page_links:
                    kind = link.get("kind")
                    _from = link.get("from")
                    bbox = tuple(_from) if _from is not None else None
                    if kind == fitz.LINK_GOTO:
                        target = link.get("page")
                        if isinstance(target, int):
                            links.append(
                                LinkInfo(
                                    type="internal",
                                    source_page=page_num + 1,
                                    target_page=target + 1,
                                    bbox=bbox,
                                )
                            )
                        elif isinstance(target, str):
                            resolved = doc.resolve_link(target)
                            links.append(
                                LinkInfo(
                                    type="internal",
                                    source_page=page_num + 1,
                                    target_page=int(resolved[0]) + 1,
                                    named_dest=target,
                                    bbox=bbox,
                                )
                            )
                    elif kind == fitz.LINK_URI:
                        links.append(
                            LinkInfo(
                                type="external",
                                uri=link.get("uri"),
                                source_page=page_num + 1,
                                bbox=bbox,
                            )
                        )
                    else:
                        logger.debug("Unsupported link kind %s on page %s: %s", kind, page_num + 1, link)  # noqa: E501

            doc.close()
            return links

        return await asyncio.to_thread(_extract)


class DocxParser(DocumentParser):
    def get_supported_extensions(self) -> list[str]:
        return [".docx"]

    async def parse(self, file_path: str) -> str:
        def _parse():
            from docx import Document
            doc = Document(file_path)
            paragraphs = []

            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            return "\n\n".join(paragraphs)

        return await asyncio.to_thread(_parse)

    async def extract_links(self, file_path: str) -> list[LinkInfo]:
        def _extract():
            from docx import Document

            doc = Document(file_path)
            links: list[LinkInfo] = []

            for para in doc.paragraphs:
                try:
                    hyperlinks = getattr(para, "hyperlinks", None)
                    if hyperlinks:
                        for hl in hyperlinks:
                            rel_id = getattr(hl, "rel_id", None) or getattr(hl, "rId", None)
                            if rel_id and rel_id in doc.part.rels:
                                uri = doc.part.rels[rel_id].target_ref
                                links.append(LinkInfo(type="external", uri=uri))
                    else:
                        for run in para.runs:
                            hl = getattr(run, "hyperlink", None)
                            if hl is not None:
                                rel_id = getattr(hl, "rel_id", None) or getattr(hl, "rId", None)
                                if rel_id and rel_id in doc.part.rels:
                                    uri = doc.part.rels[rel_id].target_ref
                                    links.append(LinkInfo(type="external", uri=uri))
                except Exception as e:
                    logger.warning("Failed to extract hyperlink from paragraph: %s", e)

            return links

        return await asyncio.to_thread(_extract)


class TextParser(DocumentParser):
    def get_supported_extensions(self) -> list[str]:
        return [".txt", ".md"]

    async def parse(self, file_path: str) -> str:
        def _parse():
            with open(file_path, encoding="utf-8", errors="ignore") as f:
                return f.read()

        return await asyncio.to_thread(_parse)


class OpenAPIParser:
    def get_supported_extensions(self) -> list[str]:
        return [".yaml", ".yml", ".json"]

    async def parse(self, file_path: str) -> str:
        from .openapi_parser import OpenAPIParser as OpenAPIParserImpl
        parser = OpenAPIParserImpl()
        return await asyncio.to_thread(parser.parse_to_text, file_path)


class ParserRegistry:
    def __init__(self):
        self.parsers: list[DocumentParser] = [
            PDFParser(),
            DocxParser(),
            TextParser(),
            OpenAPIParser(),
        ]

    def get_parser(self, file_path: str) -> DocumentParser | None:
        for parser in self.parsers:
            if any(file_path.lower().endswith(ext) for ext in parser.get_supported_extensions()):
                return parser
        return None

    async def parse(self, file_path: str) -> str:
        parser = self.get_parser(file_path)
        if parser is None:
            raise ValueError(f"No parser found for file: {file_path}")
        return await parser.parse(file_path)

    async def extract_links(self, file_path: str) -> list[LinkInfo]:
        parser = self.get_parser(file_path)
        if parser is None:
            raise ValueError(f"No parser found for file: {file_path}")
        return await parser.extract_links(file_path)
